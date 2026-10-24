import unicodedata

from docx import Document
from pathlib import Path

PROJECT_DIR = Path(__file__).resolve().parent.parent
SOURCE_DIR = PROJECT_DIR / "data" / "source"
GENERATED_DIR = PROJECT_DIR / "generated"


def normalize_name(value: str) -> str:
    value = unicodedata.normalize("NFC", value)
    return value.casefold()


def list_templates() -> list[dict]:
    """Возвращает список исходных модулей DOCX"""

    SOURCE_DIR.mkdir(parents=True, exist_ok=True)

    templates = []

    for path in sorted(SOURCE_DIR.rglob("*.docx")):
        templates.append({
            "name": path.stem,
            "filename": path.name,
            "path": str(path),
        })

    return templates


def find_template(template_name: str) -> Path:
    """Находит исходный DOCX-модуль по имени"""

    templates = list_templates()

    search_name = normalize_name(template_name)

    matches = [
        item for item in templates
        if normalize_name(item["name"]) == search_name
           or normalize_name(item["filename"]) == search_name
    ]

    if not matches:
        raise FileNotFoundError(
            f"Модуль не найден: {template_name}"
        )

    if len(matches) > 1:
        raise ValueError(
            "Найдено несколько модулей с таким именем. "
            "Уточните название."
        )

    return Path(matches[0]["path"])



def get_template_info(template_name: str) -> dict:
    """Возвращает сведения об исходном модуле"""
    path = find_template(template_name)

    descriptions = {
        "Модуль_распоряжения_о_создании_рабочей_группы": {
            "purpose": "Для оформления распоряжения о создании рабочей группы.",
            "when_to_use": "Когда необходимо сформировать рабочую группу для выполнения задачи.",
        },
        "Модуль_распоряжения_о_создании_проектного_комитета": {
            "purpose": "Для оформления распоряжения о создании проектного комитета.",
            "when_to_use": "Когда необходимо создать проектный комитет для управления проектом.",
        },
        "Модуль_приказа_о_вводе_в_эксплуатацию_информационной_системы": {
            "purpose": "Для оформления приказа о вводе информационной системы в эксплуатацию.",
            "when_to_use": "Когда информационную систему необходимо официально ввести в эксплуатацию.",
        },
        "Модуль_приказа_о_приеме_и_передаче_дел_и_имущества": {
            "purpose": "Для оформления приёма и передачи дел и имущества.",
            "when_to_use": "При передаче дел и имущества между ответственными лицами.",
        },
        "Модуль_приказа_о_назначении_ответственных_за_делопроизводство": {
            "purpose": "Для оформления назначения ответственных за делопроизводство.",
            "when_to_use": "Когда требуется назначить сотрудников, отвечающих за делопроизводство.",
        },
        "Модуль_приказа_об_утверждении_и_введении_в_действие_локального": {
            "purpose": "Для оформления утверждения и введения в действие локального документа.",
            "when_to_use": "Когда необходимо утвердить локальный документ и установить дату начала его действия.",
        },
        "Модуль_распоряжения_о_проведении_внутреннего_аудита_процессов": {
            "purpose": "Для оформления распоряжения о проведении внутреннего аудита процессов.",
            "when_to_use": "Когда необходимо организовать внутренний аудит процессов.",
        },
    }

    info = descriptions.get(
        path.stem,
        {
            "purpose": "Описание назначения пока не подготовлено.",
            "when_to_use": "Уточните назначение документа по его содержимому и методическим рекомендациям.",
        },
    )

    return {
        "name": path.stem,
        "filename": path.name,
        "format": "docx",
        "status": "source",
        **info,
        "message": "Найден исходный модуль. Оригинал не изменялся.",
    }



def read_template_content(template_name: str) -> dict:
    """Читает текст исходного DOCX-модуля"""

    path = find_template(template_name)
    document = Document(path)

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    tables = []
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]
            tables.append(" | ".join(cells))

    content = paragraphs + tables

    return {
        "name": path.stem,
        "filename": path.name,
        "content": "\n".join(content),
    }


